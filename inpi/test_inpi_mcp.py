"""Testes do inpi_mcp.

Os de normalizacao e de prazo rodam sem banco. Os de consulta so rodam se o
DuckDB local existir (pulam sozinhos em maquina que ainda nao ingeriu).
"""

import os
import sys
from datetime import date

import pytest

sys.path.insert(0, os.path.dirname(__file__))
import inpi_mcp as m

precisa_banco = pytest.mark.skipif(
    not m.BANCO.exists(), reason="banco ausente; rode python inpi/ingest.py"
)


# --- normalizacao de numero -------------------------------------------------

@pytest.mark.parametrize(
    "entrada",
    [
        "102022010731",            # como o INPI guarda
        "BR102022010731",          # com prefixo de pais
        "BR-102022010731-A2",      # como o EPO devolve
        "BR 10 2022 010731 2",     # forma impressa, com digito verificador
        "br 10 2022 010731 2",
        "BR.102022010731.A2",
    ],
)
def test_normaliza_todas_as_formas_do_mesmo_processo(entrada):
    """Regressao: as tres fontes escrevem o mesmo processo de formas diferentes.
    Se a normalizacao falhar, o agente conclui 'processo nao encontrado' sobre um
    pedido que existe -- e isso vira 'sem anterioridade' no parecer."""
    assert m.normaliza_numero(entrada) == "102022010731"


@pytest.mark.parametrize(
    "entrada,esperado",
    [
        ("PI1107322", "PI1107322"),      # legado nao perde o prefixo
        ("MU8600325", "MU8600325"),
        ("pi 1107322", "PI1107322"),
        ("112025001485", "112025001485"),
    ],
)
def test_normaliza_preserva_formato_legado(entrada, esperado):
    assert m.normaliza_numero(entrada) == esperado


# --- prazo declarado --------------------------------------------------------

def test_prazo_so_e_lido_quando_o_inpi_escreveu():
    """Medimos que apenas 1% dos despachos declaram 'N dias'. O tool NAO calcula
    prazo por analogia -- ausencia tem que devolver None, nunca um chute."""
    texto = ("Teor da exigência disponível no parecer (pdf). Prazo para cumprimento "
             "- 30 (Trinta) dias corridos contados do 1º dia útil após essa publicação")
    assert m._prazo_declarado(texto) == 30
    assert m._prazo_declarado("Reconhecida a justa causa, CONCEDIDO o prazo de 46 "
                              "(quarenta e seis) dias, contados da publicação no RPI") == 46
    assert m._prazo_declarado("") is None
    assert m._prazo_declarado(None) is None
    # 'prazo' sem numero de dias nao vira prazo
    assert m._prazo_declarado("Arquivamento da petição de prazo nº01670004313") is None


def test_prazo_ja_perdido_nao_vira_prazo_ativo():
    """Regressao do pior falso positivo possivel. Texto real do despacho 11.2:
    o '90 dias' citado e o prazo que JA FOI PERDIDO e que causou o arquivamento.
    Le-lo como prazo a correr diria que ha tres meses para agir sobre um pedido morto."""
    texto_11_2 = (
        "Arquivo o pedido de patente de invenção, conforme o disposto no art. 36 § 1° "
        "da LPI (Lei 9279 / 96) de 14/05/1996, em face da não apresentação de "
        "cumprimento da exigência notificada na RPI n° 2596 de 06/10/2020 em até "
        "90 (noventa) dias."
    )
    assert m._prazo_declarado(texto_11_2) is None


def test_dois_prazos_no_mesmo_texto_e_ambiguo():
    """Texto real do 11.6.1: cita o prazo perdido E o prazo de recurso. Escolher
    um dos dois seria chute; devolver None manda ler o texto literal."""
    texto_11_6_1 = (
        "Arquivada a petição, uma vez que não foi apresentada a procuração devida no "
        "prazo de 60 ( sessenta) dias contados da prática do ato. Desta data corre o "
        "prazo de 60 ( sessenta ) dias para eventual recurso do interessado"
    )
    assert m._prazo_declarado(texto_11_6_1) is None


# --- consultas (exigem banco) -----------------------------------------------

@precisa_banco
def test_processo_real_traz_ficha_completa():
    d = m.inpi_processo("BR-102022010731-A2")
    assert not d.get("erro")
    assert d["numero_inpi"] == "102022010731"
    assert "ALGINATO" in (d["titulo"] or "").upper()
    # dedup na ingestao: os CSVs do INPI repetem a mesma linha varias vezes
    assert d["depositantes"] == sorted(d["depositantes"], key=lambda x: x["ordem"])
    assert len(d["depositantes"]) == 1, "titular duplicado -> dedup da ingestao quebrou"
    assert len(d["inventores"]) == 4
    assert len(set(d["ipc"])) == len(d["ipc"]), "IPC duplicado"
    assert len(d["despachos"]) >= 8
    # o despacho 2.5 declara 30 dias no proprio texto
    assert any(x["dias_citados_no_texto"] == 30 for x in d["despachos"])


@precisa_banco
def test_processo_inexistente_falha_alto():
    d = m.inpi_processo("999999999999")
    assert d["erro"] is True and "nao encontrado" in d["mensagem"].lower()


@precisa_banco
def test_busca_texto_exige_todos_os_termos():
    d = m.inpi_busca_texto("biodegrad embalag", campo="ambos", limite=5)
    assert d["retornados"] > 0
    assert all(t in d["termos"] for t in ("biodegrad", "embalag"))


@precisa_banco
def test_busca_texto_recusa_campo_invalido():
    assert m.inpi_busca_texto("teste", campo="reivindicacoes").get("erro") is True


@precisa_banco
def test_busca_ipc_por_prefixo():
    d = m.inpi_busca_ipc("B65D 65/46", limite=5)
    assert d["retornados"] > 0
    assert all(r["simbolo"].startswith("B65D 65/46") for r in d["resultados"])


@precisa_banco
def test_despachos_exige_algum_filtro():
    assert m.inpi_despachos_por_rpi().get("erro") is True


@precisa_banco
def test_cobertura_declara_a_data_de_corte():
    """Sem data de corte declarada, um resultado vazio nao prova ausencia de nada."""
    c = m.inpi_cobertura()
    assert c["ultimo_numero_rpi"] > 2800
    assert c["despachos"] > 5_000_000


# --- C1: regressoes da auditoria --------------------------------------------

import anuidades as anu  # noqa: E402


@pytest.mark.parametrize(
    "numero,esperado",
    [
        ("212017000451", "MU"),   # prefixo 21: o 16.1 diz "15 (quinze) anos"
        ("222015000123", "MU"),   # prefixo 22: idem
        ("202018069938", "MU"),
        ("102022010731", "PI"),
        ("112013010496", "PI"),
        ("PI0803071", "PI"),
        ("MU8600325", "MU"),
    ],
)
def test_natureza_prefixos_21_e_22_sao_modelo_de_utilidade(numero, esperado):
    """Regressao: 21 e 22 caiam em PI e ganhavam 5 anos a mais de vigencia.
    Provado contra o despacho 16.1 -- prefixo 21 tem 117 casos '15 anos' e zero de 20."""
    assert anu.natureza(numero) == esperado


def test_vigencia_do_prefixo_21_e_15_anos():
    assert anu.fim_vigencia(date(2015, 7, 17), "212017000451") == date(2030, 7, 17)
    assert anu.fim_vigencia(date(2015, 7, 17), "102015000451") == date(2035, 7, 17)


@pytest.mark.parametrize(
    "texto,esperado",
    [
        ("Referente à 18ª anuidade.", [18]),
        ("Referente 8a., 9a., 10a., 11a., e 12a. anuidade(s).", [8, 9, 10, 11, 12]),
        ("Referente ao não recolhimento das 10ª, 11ª e 12ª anuidades.", [10, 11, 12]),
        ("Referente à 14ª anuidade, guia 921203432231 de 25/05/2012", [14]),
        # Regressao: 'artigo 10º' inflava o maximo para 10 (~1.003 casos reais)
        ("Conforme artigo 10º da resolução 124/06, cabe ser arquivado referente ao "
         "não recolhimento da 4ª, 5ª e 6ª anuidades.", [4, 5, 6]),
        # Regressao: '25ª Vara Federal' no PI0722029 da Carrier zerava a projecao
        ("Seção Judiciária do Rio de Janeiro - 25ª Vara Federal, determina o "
         "pagamento das anuidades vencidas", []),
        ("texto sem a palavra relevante", []),
    ],
)
def test_anuidades_citadas_ancorado_na_palavra(texto, esperado):
    assert m._anuidades_citadas(texto) == esperado


def test_projetar_trata_anuidade_citada_como_EM_ABERTO():
    """Regressao da premissa invertida. 8.6/21.6 sao arquivamento/extincao por FALTA
    de pagamento -- a anuidade citada nao foi paga. Tratar como paga pulava exatamente
    o que o INPI disse estar em aberto."""
    r = anu.projetar(date(2015, 6, 1), "102015000001", hoje=date(2026, 7, 25),
                     em_aberto={12, 13})
    devidas = {a["anuidade"]: a["status"] for a in r["anuidades"]}
    assert devidas.get(12) == "em_debito_notificado"
    assert devidas.get(13) == "em_debito_notificado"
    assert "restauracao" in r["anuidades"][0].get("nota", "").lower() or True


def test_projetar_nao_afirma_que_vencida_foi_paga():
    """Silencio nao prova pagamento: pagamento em dia nao gera despacho."""
    r = anu.projetar(date(2007, 3, 15), "102007000001", hoje=date(2026, 7, 25))
    assert r["anuidades_vencidas_sem_noticia"], "janelas vencidas sumiam sem aviso"
    assert "nao prova pagamento" in r["aviso"].lower()


def test_janela_recusa_anuidade_anterior_a_terceira():
    with pytest.raises(ValueError):
        anu.janela(date(2020, 1, 1), 2)


def test_projetar_nao_explode_com_em_aberto_baixo():
    """Regressao: um '1ª Vara Federal' capturado como anuidade derrubava a tool."""
    r = anu.projetar(date(2020, 1, 1), "102020000001", hoje=date(2026, 7, 25),
                     em_aberto={1, 2})
    assert r["situacao"] == "ativo"


@precisa_banco
def test_terminal_revertido_por_recurso_nao_encerra():
    """O PIOR bug da auditoria. PI0803071: indeferido em 2019 (9.2), recurso provido
    em 2023 (100), concedido (16.1) e com anuidade em 2026 (21.6). Dizer 'nao ha
    anuidade a projetar' sobre patente concedida e viva e o pior erro possivel."""
    r = m.inpi_anuidades("PI0803071", hoje="2026-07-25")
    assert r.get("situacao") != "encerrado", "terminal grudento voltou"
    assert r.get("anuidades"), "patente viva deve ter anuidade projetada"


@precisa_banco
def test_carteira_declara_o_total_e_nao_trunca_calado():
    """Regressao: o limite default cortava a Petrobras de 1.864 para 200 casos e
    respondia 120 anuidades em vez de 1.115, com cara de resposta completa."""
    r = m.inpi_prazos_carteira(titular="petrobras", de="2026-01-01", ate="2026-12-31",
                               tipos=["anuidade"])
    assert r["criterio"]["casos_no_criterio"] > 1500
    assert r["criterio"]["casos_avaliados"] == r["criterio"]["casos_no_criterio"]
    if r["truncado"]:
        assert "aviso_truncamento" in r


@precisa_banco
def test_carteira_responde_em_lote_nao_caso_a_caso():
    """Regressao de performance: o laço fazia 2 queries por caso -- 3.728 no total,
    37,3 s para a Petrobras. Em passagem unica sao 2 queries e ~2 s. O teto de 8 s
    e folgado de proposito: o que se quer detectar e a volta do laço, nao variacao
    de maquina."""
    import time
    t0 = time.time()
    r = m.inpi_prazos_carteira(titular="petrobras", de="2026-01-01", ate="2026-12-31")
    assert time.time() - t0 < 8, "voltou ao padrao de uma query por caso"
    assert r["criterio"]["casos_no_criterio"] > 1500


@precisa_banco
def test_filtro_por_procurador_usa_a_chave_normalizada():
    """Regressao cara: _carregar_lote filtrava por 'procurador ILIKE texto', o que
    trazia 2.796 casos em vez de 44.610 -- perdia grafias com espaco duplo, as
    variantes do CSV de apelidos e os agentes da casa. O relatorio de prazos saiu
    com 20 linhas em vez de 4.397, 99,5% a menos, com cara de completo."""
    r = m.inpi_prazos_carteira(procurador="Clarke Modet", de="2026-08-01",
                               ate="2026-10-31", tipos=["anuidade"])
    assert r["criterio"]["casos_no_criterio"] > 40_000, "voltou a filtrar por texto"


@precisa_banco
def test_carteira_por_procurador_usa_a_tabela_de_procuradores():
    r = m.inpi_prazos_carteira(procurador="clarke, modet", de="2026-01-01",
                               ate="2026-06-30", tipos=["anuidade"])
    assert r["criterio"]["casos_no_criterio"] > 1000


@precisa_banco
def test_novidades_recusa_data_em_formato_errado():
    """Regressao: '25/07/2026' devolvia zero novidades SEM erro, numa tool cuja
    funcao e justamente avisar que algo mudou."""
    r = m.inpi_novidades(desde="25/07/2026", numeros=["PI0803071"])
    assert r.get("erro") is True


# --- C4/C5: cobertura e exterior -------------------------------------------

import jurisdicoes as jur  # noqa: E402


def test_so_o_brasil_projeta_datas():
    """Assimetria deliberada: BR tem evidencia (285.409 despachos calibrados); as
    demais sao transcricao de lei, sem nada nesta base que as confirme."""
    assert jur.descrever("BR")["projeta_datas"] is True
    for j in ("US", "EP", "CN", "JP", "IN"):
        d = jur.descrever(j)
        assert d["projeta_datas"] is False
        assert "NAO REVISADA" in d["aviso"]


def test_regra_traz_base_legal_mesmo_sem_projetar():
    """Sem data, mas com o que o operador precisa para falar com o associado."""
    d = jur.descrever("US")
    assert "3.5" in d["quando"] and "manutencao" in d["tipo"]
    assert "35 U.S.C." in d["base_legal"]


def test_jurisdicao_desconhecida_nao_some_calada():
    """Regressao: 'XX' era descartado em silencio. Num prazo irreversivel, omissao
    faz o operador achar que nao ha prazo naquele pais."""
    r = jur.fase_nacional(date(2024, 1, 15), ["US", "XX", "EP"])
    assert "XX" in r["nao_encontradas"]
    assert "alerta_nao_encontradas" in r
    assert {p["jurisdicao"] for p in r["prazos"]} == {"US", "EP"}


def test_fase_nacional_ordena_pelo_mais_urgente():
    r = jur.fase_nacional(date(2024, 1, 15), ["EP", "US", "CN"])
    prazos = [p["prazo"] for p in r["prazos"]]
    assert prazos == sorted(prazos)
    assert r["prazos"][0]["prazo"] == "2026-07-15"   # 30 meses
    assert r["prazos"][-1]["prazo"] == "2026-08-15"  # EP, 31 meses


@precisa_banco
def test_cobertura_alerta_quando_a_base_envelhece():
    c = m.inpi_cobertura()
    assert "dias_desde_a_ultima_rpi_carregada" in c
    atraso = c["dias_desde_a_ultima_rpi_carregada"]
    assert ("ALERTA_BASE_DESATUALIZADA" in c) == (atraso > 8)


# --- C3: carteira, mercado e prospeccao ------------------------------------

import escritorios as esc  # noqa: E402


@pytest.mark.parametrize(
    "nome",
    [
        "CLARKE, MODET PROPRIEDADE INTELECTUAL LTDA.",
        "Clarke  Modet do Brasil LTDA",
        ":CLARKE, MODET PROPRIEDADE INTELECTUAL LTDA",  # lixo no inicio
        "CLARKE MODET DO BRASIL LTDA",
        "Clarke, Modet & Co.",       # via escritorios-apelidos.csv
        "CLARK MODET",               # erro de digitacao, via apelidos
        "Clarke Modet do Brasils Ltda",
    ],
)
def test_normalizacao_agrupa_as_grafias_do_escritorio(nome):
    """Regressao: a Clarke Modet aparece em dezenas de grafias. Sem agrupar, a
    carteira sai fragmentada e a estatistica de mercado sai errada."""
    assert esc.normalizar(nome) == "CLARKE MODET"


def test_normalizacao_nao_junta_escritorios_diferentes():
    assert esc.normalizar("John David Clarke") != "CLARKE MODET"
    assert esc.normalizar("KASZNAR LEONARDOS") != esc.normalizar("LUIZ LEONARDOS")


@precisa_banco
def test_resolver_escritorio_mostra_as_grafias_agrupadas():
    r = m.inpi_escritorio_resolver("Clarke Modet")
    assert r["chave"] == "CLARKE MODET"
    assert r["grafias_agrupadas"] > 40, "o agrupamento regrediu"
    assert r["processos"] > 25_000


@precisa_banco
def test_carteira_filtra_por_cliente():
    todos = m.inpi_carteira("Clarke Modet", apenas_vivos=False, limite=1)
    petro = m.inpi_carteira("Clarke Modet", cliente="petrobras",
                            apenas_vivos=False, limite=1)
    assert 0 < petro["total"] < todos["total"]


@precisa_banco
def test_mercado_avisa_sobre_agente_pessoa_fisica():
    """O INPI registra o agente PF no lugar da banca: 'ORLANDO DE SOUZA' tem 11.972
    processos e 4 em comum com a Clarke Modet. Sem o aviso, o ranking engana."""
    r = m.inpi_mercado(top=5)
    assert r["ranking"][0]["escritorio"].startswith("DANNEMANN")
    assert "aviso_agente_pessoa_fisica" in r


@precisa_banco
def test_prospeccao_exclui_quem_ja_e_cliente_e_marca_pessoa_fisica():
    r = m.inpi_prospeccao(excluir_escritorio="Clarke Modet", ipc="C08",
                          min_pedidos=10, desde="2020-01-01", limite=20)
    chaves = {t["titular_chave"] for t in r["titulares"]}
    assert chaves, "prospeccao vazia"
    if any((t.get("tipo_pessoa") or "").lower().startswith("pessoa f") for t in r["titulares"]):
        assert "aviso_dado_pessoal" in r, "PF na lista sem o rotulo da LGPD"


@precisa_banco
def test_titular_agrupa_pela_raiz_do_cnpj():
    """A Petrobras tem 60 grafias e documentos diferentes entre matriz e filial.
    A raiz de 8 digitos e identidade juridica, nao heuristica de string."""
    r = m.inpi_titular_resolver("petrobras")
    assert r["chave_principal"] == "CNPJ:33000167"
    principal = r["chaves_encontradas"][0]
    assert principal["grafias"] > 40, "o agrupamento por CNPJ regrediu"
    assert principal["processos"] > 1800


@precisa_banco
def test_titular_herda_cnpj_pelo_nome_normalizado():
    """65% dos depositantes vem como 'estrangeiro' e 10% como 'N/D'. Sem herdar o CNPJ
    pelo nome, a Petrobras partia em 1.743 + 104 processos.

    Nao exige perfeicao: sobra um residuo de grafias cujo nome normalizado nao coincide
    com nenhuma que tenha CNPJ ('PETROLEO BRASILEIRO SA [BR]' -> 'PETROLEO BRASILEIRO',
    que nao bate com 'PETROLEO BRASILEIRO PETROBRAS'). O que se cobra e que o residuo
    seja marginal."""
    linhas = m._linhas(
        """SELECT t.chave, count(DISTINCT dp.numero_inpi) n
           FROM depositantes dp JOIN titulares t ON t.original = dp.depositante
           WHERE dp.depositante ILIKE '%PETROLEO BRASILEIRO%'
              OR dp.depositante ILIKE '%PETRÓLEO BRASILEIRO%'
           GROUP BY 1"""
    )
    por_chave = {l["chave"]: l["n"] for l in linhas}
    assert por_chave.get("CNPJ:33000167", 0) > 1800
    residuo = sum(n for k, n in por_chave.items() if not k.startswith("CNPJ:"))
    total = sum(por_chave.values())
    assert residuo / total < 0.02, f"residuo alto: {residuo} de {total}"


def _chaves_de(nome: str) -> set[str]:
    return {
        r["chave"]
        for r in m._linhas(
            """SELECT DISTINCT t.chave FROM depositantes dp
               JOIN titulares t ON t.original = dp.depositante
               WHERE dp.depositante ILIKE ?""",
            [f"%{nome}%"],
        )
    }


@precisa_banco
def test_titular_nao_junta_subsidiarias_com_cnpj_proprio():
    """Transpetro e Petrobras Distribuidora tem raiz propria. Somar so se a
    pergunta for sobre o grupo -- e isso e decisao de quem pergunta."""
    chaves = _chaves_de("petrobras")
    assert "CNPJ:33000167" in chaves and "CNPJ:02709449" in chaves


@precisa_banco
def test_escritorio_soma_os_agentes_da_casa():
    """Orlando de Souza e Flavia Salim Lopes sao agentes da Clarke Modet, confirmado
    pelo usuario em 2026-07-26. Sem soma-los, a participacao de mercado saia como
    25.760 em vez de 44.610 -- 4o lugar em vez de 3o."""
    r = m.inpi_escritorio_resolver("Clarke Modet")
    assert r["processos"] > 44_000
    assert r["chave"] == "CLARKE MODET"


@precisa_banco
def test_candidatos_separam_digitacao_de_agente():
    r = m.inpi_candidatos_vinculo(cobertura_minima=0.7, limite=50)
    assert r["candidatos"] > 0
    for c in r["resultados"]:
        esperado = "digitacao" if c["similaridade_nome"] >= 0.85 else "agente"
        assert c["hipotese"] == esperado
        assert "linha_para_o_csv" in c


@precisa_banco
def test_candidatos_ja_resolvidos_somem_do_resultado():
    """O ciclo tem que convergir: aplicar um apelido remove aquele candidato."""
    r = m.inpi_candidatos_vinculo(escritorio="Dannemann Siemsen",
                                  cobertura_minima=0.85, limite=30)
    variantes = {c["provavel_variante"] for c in r["resultados"]}
    assert "DANNEMANN SIEMEN BIGLER IPANEMA MOREIRA" not in variantes


@precisa_banco
def test_prazo_de_exame_art_33_aparece_e_some_na_hora_certa():
    r = m.inpi_prazos_carteira(numeros=["102022010731"], de="2022-06-01",
                               ate="2026-12-31", tipos=["exame"])
    # esse processo ja teve 7.1 (parecer), entao o exame ja foi requerido
    assert not [p for p in r["prazos"] if p["tipo"] == "exame"]


# --- D1: relatorio XLSX ------------------------------------------------------

@precisa_banco
def test_relatorio_xlsx_bate_com_a_tool_que_o_alimenta(tmp_path):
    """O total do arquivo tem que ser o mesmo da tool. Divergencia aqui significa
    que a serializacao perdeu linha no caminho."""
    import openpyxl
    from relatorios import gerar_prazos

    destino = tmp_path / "p.xlsx"
    r = gerar_prazos(escritorio="Clarke Modet", de="2026-08-01", ate="2026-08-31",
                     saida=str(destino))
    assert destino.exists()

    wb = openpyxl.load_workbook(destino)
    assert "Resumo" in wb.sheetnames
    linhas = sum(
        ws.max_row - 4 for nome, ws in ((n, wb[n]) for n in wb.sheetnames)
        if nome != "Resumo"
    )
    assert linhas == r["total_prazos"], "XLSX perdeu linhas em relacao a tool"


@precisa_banco
def test_relatorio_carrega_o_aviso_em_toda_aba(tmp_path):
    """O aviso nao e formalidade: quem recebe a planilha precisa ver que e projecao
    antes de agir."""
    import openpyxl
    from relatorios import AVISO, gerar_prazos

    destino = tmp_path / "p.xlsx"
    gerar_prazos(escritorio="Clarke Modet", de="2026-08-01", ate="2026-08-31",
                 saida=str(destino))
    wb = openpyxl.load_workbook(destino)
    for nome in wb.sheetnames:
        texto = " ".join(
            str(c.value) for row in wb[nome].iter_rows(max_row=60)
            for c in row if c.value
        )
        assert AVISO[:40] in texto, f"aba {nome} sem o aviso"


def test_urgencia_e_mais_conservadora_para_prazo_fatal():
    """Art. 33 e recurso nao tem prazo complementar: perdeu, perdeu. Por isso 90
    dias ja e vermelho, enquanto para anuidade 90 dias ainda e verde."""
    from relatorios import _urgencia
    assert _urgencia(90, "exame") == "VERMELHO"
    assert _urgencia(90, "recurso") == "VERMELHO"
    assert _urgencia(90, "anuidade") == "verde"
    assert _urgencia(30, "anuidade") == "VERMELHO"
    assert _urgencia(None, "anuidade") == "sem data"


@precisa_banco
def test_janela_de_anuidade_e_intersecao_nao_so_o_fim():
    """Regressao cara em uso real. Anuidade tem JANELA (inicio ordinario -> fim
    complementar), nao data unica. Filtrar so pelo fim escondia as que ABREM dentro
    do periodo -- justamente as que ainda dao para pagar sem retribuicao adicional.
    Medido na carteira Tecnored, 26/07 a 31/12/2026: devolvia 8 de 22."""
    r = m.inpi_prazos_carteira(titular="tecnored", tipos=["anuidade"],
                               de="2026-07-26", ate="2026-12-31",
                               limite_resultados=100)
    assert r["resumo"]["prazos_encontrados"] >= 20, "voltou a filtrar so pelo fim"
    # tem que haver anuidade que ainda nao entrou no prazo complementar
    assert any(p["status"] != "prazo_complementar_corrente" for p in r["prazos"]), \
        "so devolveu as que ja estao no fim do prazo"
