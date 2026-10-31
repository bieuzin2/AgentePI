"""CLI: python -m patentcost.cli exportar|estatisticas ...

exportar     — portfólio público provável de um titular → PatentCost_Entrada_<cliente>.xlsx, insights.json,
               conciliacao.json (se --extracao-cliente) e lacunas.json, em --saida (pasta client-data/, fora do Git).
estatisticas — estatística agregada por seção IPC (INPI) → research/stats/BR.json do PatentCost (pública).
"""
from __future__ import annotations

import argparse
import hashlib
import json
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from patentcost.export import casos_br, casos_familia, conciliar, escrever_xlsx
from patentcost.insights import montar_insights
from patentcost.stats import estatistica_por_secao

PATENTCOST = Path(r"C:\Users\Gabriel\Desktop\PatentCost")


def _ler_extracao(arq: Path) -> list[dict]:
    ws = load_workbook(arq, read_only=True, data_only=True).worksheets[0]
    linhas = list(ws.iter_rows(values_only=True))
    cab = [str(c or "").strip().lower() for c in linhas[0]]
    return [dict(zip(cab, [str(v).strip() if v is not None else None for v in r])) for r in linhas[1:] if any(r)]


def exportar(titular: str, cliente: str, saida: Path, extracao: Path | None = None, sem_familia: bool = False) -> dict:
    from inpi.consulta import inpi_cobertura, inpi_titular_resolver

    res = inpi_titular_resolver(titular)
    chaves = [c["chave"] for c in res["chaves_encontradas"]]
    br = casos_br(chaves, cliente)
    ext = [] if sem_familia else casos_familia(br)
    casos = br + ext
    cob = inpi_cobertura()
    lacunas = ["Titularidade pública não comprova pagamento de anuidades nem quem gere o caso.",
               "Membros estrangeiros vêm da família INPADOC (EPO OPS): status inferido pelo kind code; validações EP não conferidas.",
               "Pedidos depositados há menos de 18 meses podem estar em sigilo e não aparecer."]
    if sem_familia:
        lacunas.append("Família EPO não consultada nesta execução (--sem-familia).")
    saida.mkdir(parents=True, exist_ok=True)
    escrever_xlsx(casos, PATENTCOST / "dist" / "PatentCost_Modelo_Entrada.xlsx", saida / f"PatentCost_Entrada_{cliente}.xlsx")
    fontes = [{"name": "INPI dados abertos", "coverage": f"última RPI {cob.get('ultima_rpi') or cob.get('rpi_mais_recente') or '?'}; chaves {', '.join(chaves)}"},
              {"name": "EPO OPS (família INPADOC)", "coverage": "não consultado" if sem_familia else f"{len(ext)} membros estrangeiros"}]
    (saida / "insights.json").write_text(json.dumps(montar_insights(casos, cliente, fontes, lacunas), ensure_ascii=False, indent=1), encoding="utf-8")
    (saida / "lacunas.json").write_text(json.dumps({"lacunas": lacunas, "resolucao_titular": res}, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    resumo = {"casos_br": len(br), "casos_exterior": len(ext), "chaves": chaves}
    if extracao:
        conc = conciliar(casos, _ler_extracao(extracao))
        (saida / "conciliacao.json").write_text(json.dumps(conc, ensure_ascii=False, indent=1), encoding="utf-8")
        resumo["conciliacao"] = {k: len(v) for k, v in conc.items()}
    return resumo


def estatisticas(de: str, ate: str, destino: Path) -> dict:
    from inpi.consulta import inpi_cobertura

    rows = estatistica_por_secao(de, ate)
    corpo = {"format": "patentcost-stats-1", "version": f"BR-{date.today().isoformat()}", "rows": rows,
             "metodologia": "Pedidos depositados no período com decisão final publicada; exigência = despachos 6.1/2.5/7.1; seção = IPC principal (ordem 1).",
             "consulta": date.today().isoformat(), "cobertura": inpi_cobertura()}
    corpo["sha256"] = hashlib.sha256(json.dumps(rows, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text(json.dumps(corpo, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    return {"linhas": len(rows), "destino": str(destino)}


def main(argv=None):
    ap = argparse.ArgumentParser(prog="patentcost")
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("exportar")
    e.add_argument("--titular", required=True); e.add_argument("--cliente", required=True); e.add_argument("--saida", required=True, type=Path)
    e.add_argument("--extracao-cliente", type=Path); e.add_argument("--sem-familia", action="store_true")
    s = sub.add_parser("estatisticas")
    s.add_argument("--de", default="2008-01-01"); s.add_argument("--ate", default="2016-12-31")
    s.add_argument("--destino", type=Path, default=PATENTCOST / "research" / "stats" / "BR.json")
    a = ap.parse_args(argv)
    r = exportar(a.titular, a.cliente, a.saida, a.extracao_cliente, a.sem_familia) if a.cmd == "exportar" else estatisticas(a.de, a.ate, a.destino)
    print(json.dumps(r, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
