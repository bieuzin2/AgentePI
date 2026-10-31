"""Exporta o portfólio público PROVÁVEL de um titular para o contrato de entrada do PatentCost.

Titularidade em base pública não prova pagamento nem quem gere o caso: cada caso leva `fonte` e
`confianca` ('confirmado' = dado oficial do INPI; 'inferido' = membro de família INPADOC).
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from inpi.familia_epo import familia

# Status reaproveita a classificação do próprio AgentePI (códigos terminais com reversões e anulações 21.8),
# calibrada contra a base real: inpi.prazos._estado_de + inpi.panorama_carteira._classifica.
_MAPA = {"concedida_vigente": "concedida", "em_andamento": "pendente", "extinta": "extinta", "indeferida": "abandonada", "arquivada": "abandonada"}
_CONCESSAO = {"9.1", "16.1"}


def data_concessao(despachos: list[dict]) -> str | None:
    """Concessão = expedição da carta-patente (16.1, LPI art. 38); deferimento (9.1) só se a 16.1 ainda não saiu."""
    vivos = [d for d in despachos if not d.get("anulado_em") and str(d.get("codigo_despacho")) in _CONCESSAO]
    for cod in ("16.1", "9.1"):
        datas = sorted(str(d["data_rpi"]) for d in vivos if str(d.get("codigo_despacho")) == cod)
        if datas:
            return data_iso(datas[0])
    return None
# Membro estrangeiro não concedido cuja última publicação é mais antiga que isto fica fora da planilha
# (status desconhecido sem o status legal do OPS) e entra como lacuna.
ANOS_PUBLICACAO_VIVA = 5


def modalidade_de(numero: str) -> str | None:
    """Natureza pelo prefixo do número INPI: 10/11/12/13, PI, PP, C1/C2 → PI; 20/21/22, MU → MU;
    MI e 3x (desenho industrial) → None (não é patente; fica fora da exportação)."""
    n = (numero or "").upper().replace(" ", "").removeprefix("BR")
    if n.startswith("MI") or n[:1] == "3":
        return None
    return "MU" if n.startswith("MU") or n[:1] == "2" else "PI"


def secao_ipc(ipc: str | None) -> str | None:
    s = (ipc or "").strip()
    return s[0].upper() if s[:1].isalpha() else None


def data_iso(v: str | None) -> str | None:
    s = str(v or "").strip()
    if re.fullmatch(r"\d{8}", s):
        return f"{s[:4]}-{s[4:6]}-{s[6:]}"
    return s[:10] if re.match(r"\d{4}-\d{2}-\d{2}", s) else None


def classificar_status(numero: str, despachos: list[dict], data_deposito=None) -> tuple[str, str | None]:
    """(status PatentCost, data de concessão) pelo estado do AgentePI; despachos anulados não contam."""
    from inpi.panorama_carteira import _classifica
    from inpi.prazos import _estado_de

    est = _estado_de(numero, [dict(d) for d in despachos], data_deposito)
    return _MAPA[_classifica(est)], data_concessao(est["despachos"])


OBS = "Portfólio público provável; titularidade não prova pagamento nem gestão do caso."


def casos_br(chaves: list[str], cliente: str) -> list[dict]:
    """Casos BR de todas as chaves canônicas confirmadas do titular. Consulta a base do INPI (local ou MotherDuck)."""
    from inpi.consulta import _linhas
    from inpi.panorama_carteira import _classifica
    from inpi.prazos import _carregar_lote

    nums = [r["numero_inpi"] for r in _linhas(
        "SELECT DISTINCT dp.numero_inpi FROM depositantes dp JOIN titulares t ON t.original = dp.depositante WHERE list_contains(?, t.chave)", [chaves])]
    estados, _ = _carregar_lote(numeros=nums)
    bib = {r["numero_inpi"]: r for r in _linhas(
        """SELECT b.numero_inpi, min(b.data_deposito) dep, any_value(b.numero_pct) pct FROM biblio b
           WHERE list_contains(?, b.numero_inpi) GROUP BY 1""", [nums])}
    ipc = {r["numero_inpi"]: r["simbolo"] for r in _linhas(
        """SELECT numero_inpi, simbolo FROM ipc WHERE list_contains(?, numero_inpi)
           QUALIFY row_number() OVER (PARTITION BY numero_inpi ORDER BY ordem) = 1""", [nums])}
    tit = {r["numero_inpi"]: r["titulo"] for r in _linhas("SELECT numero_inpi, any_value(titulo) titulo FROM conteudo WHERE list_contains(?, numero_inpi) GROUP BY 1", [nums])}
    dep = {r["numero_inpi"]: r["d"] for r in _linhas("SELECT numero_inpi, string_agg(DISTINCT depositante, '; ') d FROM depositantes WHERE list_contains(?, numero_inpi) GROUP BY 1", [nums])}
    out = []
    for n in nums:
        if modalidade_de(n) is None:
            continue
        est = estados.get(n) or {"despachos": [], "encerrado_em": None, "despacho_terminal": None}
        b = bib.get(n, {})
        out.append({
            "id_caso": f"BR-{n}", "cliente": cliente, "jurisdicao": "BR", "modalidade": modalidade_de(n), "tipo_registro": "existente",
            "numero_pedido": n, "titulo": tit.get(n), "titular": dep.get(n), "rota": "pct_fase_nacional" if b.get("pct") else "direta",
            "status": _MAPA[_classifica(est)], "data_status": data_iso(str(est.get("encerrado_em") or "")),
            "data_deposito": data_iso(str(b.get("dep") or "")), "data_concessao": data_concessao(est["despachos"]),
            "ipc_cpc": ipc.get(n), "categoria_tecnologica": secao_ipc(ipc.get(n)), "categoria_requerente": "grande", "familia": n,
            "fonte": "INPI dados abertos", "confianca": "confirmado" if b.get("dep") else "inferido", "observacoes": OBS,
        })
    return out


def casos_familia(casos: list[dict], hoje: str | None = None) -> tuple[list[dict], dict]:
    """Membros estrangeiros da família INPADOC (EPO OPS, cacheado em disco) como casos inferidos.

    Agrupa pela família INPADOC (familia_id), deduplica membros por (país, nº do pedido) mantendo a publicação
    concedida, ignora WO (fase internacional) e deixa fora os não concedidos com publicação antiga.
    Devolve (casos, relatório)."""
    from datetime import date

    limite = int((hoje or date.today().isoformat())[:4]) - ANOS_PUBLICACAO_VIVA
    melhor: dict[tuple[str, str], dict] = {}
    rel = {"familias_consultadas": 0, "sem_familia": 0, "wo_ignorados": 0, "antigos_nao_concedidos": 0, "duplicados": 0}
    for c in casos:
        fam = familia(c["numero_pedido"])
        rel["familias_consultadas"] += 1
        if fam.get("erro") or not fam.get("membros"):
            rel["sem_familia"] += 1
            continue
        fid = str(fam.get("familia_id") or c["familia"])
        c["familia"] = fid
        for m in fam["membros"]:
            pais, pub = m.get("pais"), m.get("numero")
            if not pais or not pub or pais == "BR":
                continue
            if pais == "WO":
                rel["wo_ignorados"] += 1
                continue
            pedido = m.get("numero_pedido") or pub
            k = (pais, pedido)
            atual = melhor.get(k)
            if atual is not None:
                rel["duplicados"] += 1
                if atual["_concedida"] or not m.get("concedida"):
                    continue
            melhor[k] = {**m, "_concedida": bool(m.get("concedida")), "_caso": c, "_fid": fid}
    out = []
    for (pais, pedido), m in melhor.items():
        c = m["_caso"]
        grant = data_iso(m.get("data_publicacao")) if m["_concedida"] else None
        pub = data_iso(m.get("data_publicacao"))
        if not grant and (not pub or int(pub[:4]) < limite):
            rel["antigos_nao_concedidos"] += 1
            continue
        out.append({
            "id_caso": f"{pais}-{pedido}", "cliente": c["cliente"], "jurisdicao": pais, "modalidade": c.get("modalidade", "PI"),
            "tipo_registro": "existente", "numero_pedido": pedido, "numero_patente": m.get("numero") if grant else None,
            "titulo": m.get("titulo") or c.get("titulo"), "status": "concedida" if grant else "pendente",
            "data_deposito": data_iso(m.get("data_deposito")), "data_concessao": grant, "ipc_cpc": c.get("ipc_cpc"),
            "categoria_tecnologica": c.get("categoria_tecnologica"), "categoria_requerente": "grande", "familia": m["_fid"],
            "rota": "regional" if pais in ("EP", "EA", "AP", "OA") else "pct_fase_nacional" if c.get("rota") == "pct_fase_nacional" else "direta",
            "fonte": "EPO OPS INPADOC", "confianca": "inferido",
            "observacoes": OBS + " Status inferido pelo kind code da publicação; validações EP e anuidades pagas não verificadas.",
        })
    return out, rel


CAMPOS = ["id_caso", "cliente", "jurisdicao", "modalidade", "tipo_registro", "numero_pedido", "numero_patente", "familia", "titulo", "titular",
          "rota", "status", "data_status", "data_deposito", "data_concessao", "ipc_cpc", "categoria_tecnologica", "categoria_requerente", "observacoes"]


# Nº de reivindicações não está nos dados abertos do INPI: cenários de sensibilidade nomeados (o "base" fica sem
# premissa e o adicional por reivindicação excedente segue pendente até haver contagem real ou escolha do usuário).
CENARIOS_REIVINDICACOES = (10, 15, 20, 30)


def premissas_reivindicacoes(valores=CENARIOS_REIVINDICACOES) -> list[dict]:
    return [{"cenario": f"reiv_{n}", "jurisdicao": "BR", "modalidade": "PI", "categoria_tecnologica": "*", "parametro": "reivindicacoes_padrao",
             "valor": n, "origem": "ajuste", "justificativa": f"Sensibilidade: {n} reivindicações por pedido sem contagem real (não consta dos dados abertos do INPI)."}
            for n in valores]


def escrever_xlsx(casos: list[dict], modelo: Path, saida: Path, premissas: list[dict] | None = None) -> Path:
    """Preenche as abas Casos e Premissas do modelo oficial do PatentCost (cabeçalho = chaves dos campos)."""
    wb = load_workbook(modelo)
    # Comentários de ajuda do modelo regravados pelo openpyxl quebram o leitor ExcelJS do PatentCost
    # (teste de contrato tests/agentepi-contract.test.ts): o arquivo de dados sai sem eles.
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                cell.comment = None
    ws = wb["Casos"]
    cab = [c.value for c in ws[1]]
    col = {k: cab.index(k) + 1 for k in CAMPOS if k in cab}
    for i, c in enumerate(casos, 2):
        for k, j in col.items():
            v = c.get(k)
            if k == "observacoes":
                v = f"[{c.get('fonte', '?')}|{c.get('confianca', '?')}] {v or ''}".strip()
            ws.cell(i, j).value = v
    wp = wb["Premissas"]
    cab = [c.value for c in wp[1]]
    for i, p in enumerate(premissas or [], 2):
        for k, v in p.items():
            wp.cell(i, cab.index(k) + 1).value = v
    saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(saida)
    return saida


def conciliar(publico: list[dict], cliente: list[dict], campos=("status", "data_deposito", "data_concessao", "jurisdicao")) -> dict:
    """Compara a extração pública com a do cliente por numero_pedido; nada é sobrescrito."""
    p = {c["numero_pedido"]: c for c in publico}
    k = {c["numero_pedido"]: c for c in cliente}
    div = [{"numero_pedido": n, "campo": f, "publico": p[n].get(f), "cliente": k[n].get(f)}
           for n in sorted(p.keys() & k.keys()) for f in campos if k[n].get(f) not in (None, "") and p[n].get(f) != k[n].get(f)]
    return {"apenas_publico": sorted(p.keys() - k.keys()), "apenas_cliente": sorted(k.keys() - p.keys()), "divergencias": div}
